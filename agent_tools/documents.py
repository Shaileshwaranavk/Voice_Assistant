"""DOCUMENT tools: read, search, extract facts, compare, create - all on this PC.

Only `summarize_document` talks to Gemini (to write the summary). Reading uses local libraries; scans and
images use Windows' built-in OCR.
"""

import csv
import difflib
import math
import os
import re
import tempfile
from collections import Counter
from pathlib import Path

import llm
from . import tool, remember, CTX
from .files import resolve, known_folder

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tif", ".tiff"}
TEXT_EXT = {".txt", ".md", ".log", ".json", ".xml", ".html", ".htm", ".py", ".js", ".ini", ".yaml", ".yml"}
DOC_EXT = (".pdf", ".docx", ".doc", ".xlsx", ".xls", ".pptx", ".ppt", ".csv", ".txt", ".md", ".rtf")


def doc_path(p):
    """Resolve a user's file reference: a real path, or 'it' / 'that' / '' = the file we last worked on."""
    p = (p or "").strip()
    if p and p.lower() not in ("it", "that", "this", "last", "the file", "the document"):
        path = resolve(p)
        if path and path.exists():
            return path
    last = CTX.get("last_file")
    return Path(last) if last and Path(last).exists() else None


# --------------------------------------------------------------------------
# reading
# --------------------------------------------------------------------------
def _ocr_pdf(path, max_pages=8):
    try:
        import pymupdf
    except ImportError:
        return ""
    from .ocr import ocr_files
    images = []
    with pymupdf.open(str(path)) as doc:
        for i, page in enumerate(doc):
            if i >= max_pages:
                break
            img = os.path.join(tempfile.gettempdir(), f"veera_pdf_{i}.png")
            page.get_pixmap(dpi=170).save(img)
            images.append(img)
    pages = ocr_files(images)
    for img in images:
        try:
            os.remove(img)
        except OSError:
            pass
    return "\n\n".join(f"[page {i + 1}] {t}" for i, t in enumerate(pages) if t)


def extract_text(path):
    """(text, kind) for any supported file. Raises ValueError for unsupported ones."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        text = "\n".join((pg.extract_text() or "") for pg in reader.pages).strip()
        kind = f"PDF, {len(reader.pages)} pages"
        if len(text) < 15:                                  # a scan: no real text layer
            text = _ocr_pdf(path)
            kind += ", scanned (read with OCR)"
        return text, kind
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts), "Word document"
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"== sheet: {ws.title} ==")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 100:
                    parts.append("... (more rows)")
                    break
                if any(c is not None for c in row):
                    parts.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(parts), "Excel workbook"
    if ext == ".pptx":
        from pptx import Presentation
        parts = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            parts.append(f"== slide {i} ==")
            for shp in slide.shapes:
                if shp.has_text_frame and shp.text_frame.text.strip():
                    parts.append(shp.text_frame.text)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                parts.append("(notes) " + slide.notes_slide.notes_text_frame.text)
        return "\n".join(parts), "PowerPoint"
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8", errors="replace") as f:
            rows = [" | ".join(r) for _, r in zip(range(100), csv.reader(f))]
        return "\n".join(rows), "CSV"
    if ext in TEXT_EXT:
        return path.read_text(encoding="utf-8", errors="replace"), "text file"
    if ext in IMAGE_EXT:
        from .ocr import ocr_files
        return ocr_files([path])[0], "image (read with OCR)"
    if ext in (".doc", ".xls", ".ppt"):
        raise ValueError("That's an old Office format - open it in Office and save as .docx / .xlsx / .pptx.")
    raise ValueError(f"I don't know how to read {ext or 'that'} files.")


def text_of(path_arg):
    """(path, text, kind) - raises ValueError with a speakable message when it can't be read."""
    p = doc_path(path_arg)
    if not p:
        raise ValueError("I can't find that file.")
    remember(last_file=str(p), last_document=str(p))
    text, kind = extract_text(p)
    if not text.strip():
        raise ValueError(f"I couldn't find any readable text in {p.name}.")
    return p, text, kind


@tool("documents")
def read_document(path: str = "", max_chars: int = 12000) -> str:
    """Read the text of a document: PDF, Word, Excel, PowerPoint, CSV, text file, or an image / scan (OCR).
    path empty or 'it' = the file we were just working with."""
    try:
        p, text, kind = text_of(path)
    except Exception as e:
        return str(e)
    cut = "" if len(text) <= max_chars else f"\n... (truncated, {len(text)} characters in total)"
    return f"[{p.name} - {kind}]\n{text[:max_chars]}{cut}"


@tool("documents")
def summarize_document(path: str = "", style: str = "quick") -> str:
    """Summarise a document (uses Gemini - the only thing that does). style: quick | detailed | executive | action_items."""
    try:
        p, text, kind = text_of(path)
    except Exception as e:
        return str(e)
    return llm.summarize(text, style, subject=f"{p.name} ({kind})")


# --------------------------------------------------------------------------
# facts and questions, answered locally
# --------------------------------------------------------------------------
_MONTHS = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
DATE_RE = re.compile(
    rf"\b\d{{1,2}}(?:st|nd|rd|th)?[\s\-/]+{_MONTHS}\.?,?[\s\-/]+\d{{2,4}}\b|"      # 15 October 2026
    rf"\b\d{{1,2}}(?:st|nd|rd|th)?\s+(?:of\s+)?{_MONTHS}\b|"                          # 15 October
    rf"\b{_MONTHS}\.?\s+\d{{1,2}}(?:st|nd|rd|th)?(?:,?\s+\d{{2,4}})?\b|"              # October 15, 2026
    r"\b\d{1,2}[/.\-]\d{1,2}[/.\-]\d{2,4}\b|\b\d{4}-\d{2}-\d{2}\b", re.I)
MONEY_RE = re.compile(
    r"(?:₹|rs\.?|inr|\$|usd|€|eur|£|gbp)\s?\d[\d,]*(?:\.\d+)?|"
    r"\b\d[\d,]*(?:\.\d+)?\s?(?:rupees|dollars|euros|pounds|usd|inr|eur)\b", re.I)
EMAIL_RE = re.compile(r"\b[\w.+-]+@[\w-]+(?:\.[\w-]+)+\b")
PHONE_RE = re.compile(r"(?<!\d)(?:\+?\d{1,3}[\s\-]?)?(?:\(?\d{2,4}\)?[\s\-]?)?\d{3,4}[\s\-]?\d{3,4}(?!\d)")
URL_RE = re.compile(r"https?://[^\s)>\"]+|www\.[^\s)>\"]+")
ID_RE = re.compile(r"\b(invoice|order|reference|ref|ticket|po|booking|account|policy|receipt|case)\s*"
                   r"(?:no\.?|number|num|id|#|:)?\s*[:#]?\s*([A-Z0-9][A-Z0-9\-/]{2,})", re.I)
WHO_RE = re.compile(r"\b(customer|client|bill(?:ed)? to|attn|attention|dear|prepared for|to)\s*[:\-]?\s*"
                    r"([A-Z][\w.&'-]*(?:\s+[A-Z][\w.&'-]*){0,3})")
DEADLINE_WORDS = re.compile(r"\b(deadline|due|by|before|until|no later than|expires?|last date|submit|respond)\b", re.I)
TOTAL_WORDS = re.compile(r"\b(total|amount due|grand total|balance|payable|sum|net)\b", re.I)
ACTION_WORDS = re.compile(r"\b(please|must|need to|needs to|should|required|kindly|action|todo|to do|follow up|"
                          r"send|submit|review|confirm|complete|reply|respond)\b", re.I)

_STOP = set("the a an and or of to in on for with is are was were be been it this that these those at by from as "
            "what which who whom when where how much many do does did me my i you your we our they their about "
            "tell find give show".split())


def _sentences(text):
    parts = re.split(r"(?<=[.!?])\s+|\n+", text)
    return [s.strip() for s in parts if len(s.strip()) > 3]


def facts(text, what):
    """Pull one kind of fact out of text, locally. Returns a list of strings."""
    what = what.lower()
    sents = _sentences(text)
    if what in ("deadline", "due date", "due", "deadlines"):
        hits = [s for s in sents if DATE_RE.search(s) and DEADLINE_WORDS.search(s)]
        return hits or [s for s in sents if DATE_RE.search(s)][:3]
    if what in ("date", "dates"):
        return list(dict.fromkeys(m.group().strip() for m in DATE_RE.finditer(text)))
    if what in ("total", "amount", "amounts", "price", "cost", "money", "balance"):
        hits = [s for s in sents if TOTAL_WORDS.search(s) and MONEY_RE.search(s)]
        return hits or list(dict.fromkeys(m.group().strip() for m in MONEY_RE.finditer(text)))
    if what in ("email", "emails", "email address", "email addresses"):
        return list(dict.fromkeys(EMAIL_RE.findall(text)))
    if what in ("phone", "phones", "phone number", "phone numbers", "contact"):
        found = [m.group().strip() for m in PHONE_RE.finditer(text) if len(re.sub(r"\D", "", m.group())) >= 8]
        return list(dict.fromkeys(found + EMAIL_RE.findall(text)))
    if what in ("link", "links", "url", "urls"):
        return list(dict.fromkeys(URL_RE.findall(text)))
    if what in ("invoice number", "order number", "number", "id", "reference", "reference number"):
        return [f"{m.group(1)} {m.group(2)}" for m in ID_RE.finditer(text)]
    if what in ("customer", "client", "name", "names", "who"):
        return list(dict.fromkeys(m.group(2).strip() for m in WHO_RE.finditer(text)))
    if what in ("action items", "tasks", "action", "to do", "requests"):
        return [s for s in sents if ACTION_WORDS.search(s)][:12]
    return []


def best_sentences(text, question, n=3):
    """The sentences that best match the words of a question (a small TF-IDF ranking). Local."""
    sents = _sentences(text)
    if not sents:
        return []
    q = [w for w in re.findall(r"[a-z0-9]+", question.lower()) if w not in _STOP and len(w) > 2]
    if not q:
        return sents[:n]
    df = Counter()
    toks = []
    for s in sents:
        t = set(re.findall(r"[a-z0-9]+", s.lower()))
        toks.append(t)
        df.update(t)
    scored = []
    for i, t in enumerate(toks):
        score = sum(math.log(1 + len(sents) / df[w]) for w in q if w in t)
        if score:
            scored.append((score, i))
    scored.sort(reverse=True)
    return [sents[i] for _, i in sorted(scored[:n], key=lambda x: x[1])]


@tool("documents")
def find_in_document(what: str, path: str = "") -> str:
    """Find specific information in a document, locally: deadline, dates, total / amounts, emails, phone numbers,
    links, invoice or order number, customer name, action items - or any question in plain words."""
    try:
        p, text, kind = text_of(path)
    except Exception as e:
        return str(e)
    found = facts(text, what)
    if not found:
        found = best_sentences(text, what)
        if not found:
            return f"I couldn't find anything about '{what}' in {p.name}."
    return f"In {p.name}: " + " ... ".join(found[:6])


@tool("documents")
def compare_documents(path_a: str, path_b: str) -> str:
    """Compare two documents locally: how similar they are, and the lines that exist in only one of them."""
    try:
        pa, ta, _ = text_of(path_a)
        pb, tb, _ = text_of(path_b)
    except Exception as e:
        return str(e)
    la, lb = [l.strip() for l in ta.splitlines() if l.strip()], [l.strip() for l in tb.splitlines() if l.strip()]
    sim = difflib.SequenceMatcher(None, ta, tb).ratio()
    only_a = [l for l in la if l not in set(lb)][:5]
    only_b = [l for l in lb if l not in set(la)][:5]
    out = [f"{pa.name} and {pb.name} are about {round(sim * 100)} percent alike."]
    if only_a:
        out.append(f"Only in {pa.name}: " + " | ".join(only_a))
    if only_b:
        out.append(f"Only in {pb.name}: " + " | ".join(only_b))
    if not only_a and not only_b:
        out.append("Their lines are identical.")
    return " ".join(out)


@tool("documents")
def read_image_text(path: str = "") -> str:
    """Read the text in an image / screenshot / photo of a document (local OCR)."""
    p = doc_path(path)
    if not p or p.suffix.lower() not in IMAGE_EXT:
        return "I need an image file for that."
    remember(last_file=str(p))
    try:
        from .ocr import ocr_files
        return ocr_files([p])[0] or "I couldn't see any text in that image."
    except Exception as e:
        return str(e)


# --------------------------------------------------------------------------
# creating and finding
# --------------------------------------------------------------------------
def _unique(path):
    if not path.exists():
        return path
    for i in range(2, 500):
        q = path.with_name(f"{path.stem} ({i}){path.suffix}")
        if not q.exists():
            return q
    return path


@tool("documents")
def create_document(kind: str, name: str = "", folder: str = "desktop", text: str = "") -> str:
    """Create a REAL document: kind = word | excel | powerpoint | text | csv. name without extension.
    text = optional starting content. Never overwrites an existing file."""
    k = kind.lower()
    ext = {"word": ".docx", "docx": ".docx", "document": ".docx", "excel": ".xlsx", "xlsx": ".xlsx",
           "spreadsheet": ".xlsx", "powerpoint": ".pptx", "pptx": ".pptx", "presentation": ".pptx",
           "slides": ".pptx", "text": ".txt", "txt": ".txt", "notepad": ".txt", "note": ".txt",
           "csv": ".csv", "python": ".py", "py": ".py", "json": ".json", "html": ".html", "markdown": ".md",
           "batch": ".bat", "javascript": ".js", "css": ".css"}.get(k)
    if not ext:
        return "I can create word, excel, powerpoint, text, csv, python, json, html or markdown files."
    base = resolve(folder) or known_folder("desktop")
    base.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r'[\\/:*?"<>|]', "", name).strip() or {"docx": "New Document", "xlsx": "New Spreadsheet",
                                                         "pptx": "New Presentation", "py": "script", "json": "data", "html": "page", "md": "notes"}.get(ext[1:], "New File")
    path = _unique(base / (stem + ext))
    if ext == ".docx":
        import docx
        d = docx.Document()
        if text:
            d.add_paragraph(text)
        d.save(str(path))
    elif ext == ".xlsx":
        import openpyxl
        wb = openpyxl.Workbook()
        if text:
            wb.active["A1"] = text
        wb.save(str(path))
    elif ext == ".pptx":
        from pptx import Presentation
        pr = Presentation()
        slide = pr.slides.add_slide(pr.slide_layouts[0])
        slide.shapes.title.text = text or stem
        pr.save(str(path))
    else:
        path.write_text(text, encoding="utf-8")
    remember(last_file=str(path), last_folder=str(path.parent))
    return f"Created {path}"


@tool("documents")
def latest_document(folder: str = "downloads", extension: str = "") -> str:
    """The newest document in a folder (default Downloads); extension like 'pdf' to filter. Returns its path."""
    root = resolve(folder)
    if not root or not root.exists():
        return "That folder doesn't exist."
    kinds = {f".{extension.lower().lstrip('.')}"} if extension else set(DOC_EXT) | IMAGE_EXT
    files = [f for f in root.iterdir() if f.is_file() and f.suffix.lower() in kinds and not f.name.startswith("~$")]
    if not files:
        return "No documents there."
    newest = max(files, key=lambda f: f.stat().st_mtime)
    remember(last_file=str(newest), last_document=str(newest))
    return str(newest)
